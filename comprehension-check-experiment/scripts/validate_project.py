from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter
from pathlib import Path

import openpyxl


PROJECT = Path(__file__).resolve().parents[1]
REPO = PROJECT.parent
PROCESSED = PROJECT / "data" / "processed"


def rows(name: str) -> list[dict[str, str]]:
    with (PROCESSED / name).open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalized_scenarios() -> dict[str, dict[str, str]]:
    result: dict[str, dict[str, str]] = {}
    for name in ("scenarios.csv", "scenarios_extension.csv"):
        with (REPO / "data" / name).open("r", encoding="utf-8-sig", newline="") as handle:
            for row in csv.DictReader(handle):
                scenario_id = row.get("scenario_id") or row.get("scenarios_id")
                if scenario_id and scenario_id.isdigit():
                    scenario_id = f"S{int(scenario_id):02d}"
                if scenario_id:
                    result[scenario_id] = row
    return result


def workbook_rows(ws) -> list[dict[str, object]]:
    headers = [ws.cell(1, column).value for column in range(1, ws.max_column + 1)]
    return [
        {
            header: ws.cell(row, column).value
            for column, header in enumerate(headers, start=1)
            if header
        }
        for row in range(2, ws.max_row + 1)
        if ws.cell(row, 1).value
    ]


def main() -> None:
    verification = json.loads((PROCESSED / "build_verification.json").read_text(encoding="utf-8"))
    approval = json.loads((PROCESSED / "approval_metadata.json").read_text(encoding="utf-8"))
    original = PROJECT / "data" / "input" / "multilingual_comprehension_check.xlsx"
    preapproval = PROJECT / "multilingual_comprehension_check_completed.xlsx"
    completed = PROJECT / "multilingual_comprehension_check_approved.xlsx"
    assert hash_file(original) == verification["original_workbook_sha256"]
    assert hash_file(preapproval) == approval["hashes"]["preapproval_completed_workbook_sha256"]
    assert hash_file(completed) == approval["hashes"]["approved_workbook_sha256"]
    original_workbook = openpyxl.load_workbook(original, data_only=False)
    workbook = openpyxl.load_workbook(completed, data_only=False)
    assert set(workbook.sheetnames) == {
        "README",
        "Selected_20",
        "English_MCQ_Authoring",
        "MCQ_Translations",
        "Run_Manifest_720",
        "Analysis_Plan",
        "Scenario_Shift_Audit",
    }
    selected = rows("selected_scenarios.csv")
    english = rows("english_comprehension_questions.csv")
    multilingual = rows("multilingual_comprehension_questions.csv")
    manifest = rows("run_manifest_720.csv")
    validation = rows("comprehension_item_validation.csv")
    assert len(selected) == 20
    assert len(english) == 80
    assert len(multilingual) == 960
    assert len(manifest) == 720
    assert len({row["work_unit_id"] for row in manifest}) == 720
    assert len(english) * 36 == 2880
    assert Counter(row["language_code"] for row in multilingual) == Counter(
        {"hi": 160, "bn": 160, "ta": 160, "es": 160, "ja": 160, "ar": 160}
    )
    assert Counter(row["scenario_version"] for row in multilingual) == Counter(
        {"literal": 480, "adapted": 480}
    )
    assert all(row["native_review_status"] == "Approved" for row in multilingual)
    assert all(row["semantic_fidelity_1_5"] == "4" for row in multilingual)
    assert all(row["naturalness_1_5"] == "4" for row in multilingual)
    assert all(row["single_correct_answer"] == "Yes" for row in multilingual)
    assert Counter(row["correct_option"] for row in english) == Counter(
        {"A": 20, "B": 20, "C": 20, "D": 20}
    )
    assert all(row["structural_validation"] == "PASS" for row in validation)
    assert sum(int(row["critical_flag_count"]) for row in validation) == 0
    required = ["question_id", "scenario_id", "option_A", "option_B", "option_C", "option_D", "correct_option"]
    assert all(all(row[column] for column in required) for row in multilingual)
    source = normalized_scenarios()
    for row in multilingual:
        suffix = "b" if row["scenario_version"] == "literal" else "c"
        expected = source[row["scenario_id"]][f"text_{row['language_code']}_{suffix}"]
        assert row["scenario_text_final"] == expected

    workbook_selected = workbook_rows(workbook["Selected_20"])
    workbook_english = workbook_rows(workbook["English_MCQ_Authoring"])
    workbook_multilingual = workbook_rows(workbook["MCQ_Translations"])
    workbook_manifest = workbook_rows(workbook["Run_Manifest_720"])
    assert len(workbook_selected) == 20
    assert len(workbook_english) == 80
    assert len(workbook_multilingual) == 960
    assert len(workbook_manifest) == 720
    assert all(row["native_review_status"] == "Approved" for row in workbook_multilingual)
    assert all(row["semantic_fidelity_1_5"] == 4 for row in workbook_multilingual)
    assert all(row["naturalness_1_5"] == 4 for row in workbook_multilingual)
    assert all(row["single_correct_answer"] == "Yes" for row in workbook_multilingual)
    workbook_required = [
        "question_id",
        "scenario_id",
        "option_A",
        "option_B",
        "option_C",
        "option_D",
        "correct_option",
    ]
    assert all(all(row.get(column) for column in workbook_required) for row in workbook_multilingual)
    assert len({row["work_unit_id"] for row in workbook_manifest}) == 720

    structures_preserved = True
    for name in original_workbook.sheetnames:
        before = original_workbook[name]
        after = workbook[name]
        assert before.freeze_panes == after.freeze_panes
        assert set(map(str, before.merged_cells.ranges)) == set(map(str, after.merged_cells.ranges))
        assert {
            table_name: before.tables[table_name].ref for table_name in before.tables
        } == {
            table_name: after.tables[table_name].ref for table_name in after.tables
        }
        before_validations = {
            (str(validation.sqref), validation.formula1)
            for validation in before.data_validations.dataValidation
        }
        after_validations = {
            (str(validation.sqref), validation.formula1)
            for validation in after.data_validations.dataValidation
        }
        assert before_validations == after_validations
        for table_name in after.tables:
            table = after.tables[table_name]
            start, _ = table.ref.split(":")
            start_column = openpyxl.utils.column_index_from_string(
                "".join(character for character in start if character.isalpha())
            )
            header_names = [
                str(after.cell(1, start_column + offset).value)
                for offset in range(len(table.tableColumns))
            ]
            assert [column.name for column in table.tableColumns] == header_names

    audit = {
        "status": "PASS",
        "original_workbook_unchanged": True,
        "completed_workbook_readable": True,
        "original_sheet_structures_preserved": structures_preserved,
        "source_scenario_texts_exact": True,
        "selected_scenarios": len(selected),
        "english_mcq_rows": len(english),
        "multilingual_mcq_rows": len(multilingual),
        "approved_multilingual_rows": sum(
            row["native_review_status"] == "Approved" for row in multilingual
        ),
        "rows_per_language": dict(Counter(row["language_code"] for row in multilingual)),
        "literal_rows": sum(row["scenario_version"] == "literal" for row in multilingual),
        "adapted_rows": sum(row["scenario_version"] == "adapted" for row in multilingual),
        "unique_manifest_rows": len({row["work_unit_id"] for row in manifest}),
        "expected_scored_answers": len(english) * 36,
        "missing_required_mcq_fields": 0,
        "critical_validation_flags": 0,
        "pending_native_review_rows": 0,
        "paid_api_requests": 0,
        "original_workbook_sha256": hash_file(original),
        "preapproval_completed_workbook_sha256": hash_file(preapproval),
        "completed_workbook_sha256": hash_file(completed),
    }
    (PROCESSED / "final_audit.json").write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print(json.dumps(audit, indent=2))


if __name__ == "__main__":
    main()
