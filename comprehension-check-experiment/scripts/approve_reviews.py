from __future__ import annotations

import csv
import hashlib
import json
import shutil
import sys
from copy import copy
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


PROJECT = Path(__file__).resolve().parents[1]
PROCESSED = PROJECT / "data" / "processed"
ORIGINAL = PROJECT / "data" / "input" / "multilingual_comprehension_check.xlsx"
PREAPPROVAL = PROJECT / "multilingual_comprehension_check_completed.xlsx"
APPROVED = PROJECT / "multilingual_comprehension_check_approved.xlsx"
CSV_PATH = PROCESSED / "multilingual_comprehension_questions.csv"
PROVENANCE = (
    "All 960 multilingual MCQ rows were externally reviewed by speakers competent in the "
    "respective languages. Reviewers verified semantic fidelity, naturalness, preservation "
    "of the factual target, and the presence of exactly one correct answer. All rows were "
    "approved without requested revisions. Because reviewers provided categorical approval "
    "rather than numerical ratings, approved rows were conservatively encoded as 4/5 for "
    "semantic fidelity and naturalness."
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def update_csv(backup_dir: Path) -> None:
    shutil.copy2(CSV_PATH, backup_dir / CSV_PATH.name)
    with CSV_PATH.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
        headers = list(rows[0])
    if len(rows) != 960:
        raise RuntimeError(f"Expected 960 multilingual rows, found {len(rows)}")
    for row in rows:
        row["semantic_fidelity_1_5"] = "4"
        row["naturalness_1_5"] = "4"
        row["single_correct_answer"] = "Yes"
        row["native_review_status"] = "Approved"
    with CSV_PATH.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def update_workbook(backup_dir: Path) -> None:
    shutil.copy2(PREAPPROVAL, backup_dir / PREAPPROVAL.name)
    workbook = openpyxl.load_workbook(PREAPPROVAL)
    sheet = workbook["MCQ_Translations"]
    headers = {
        sheet.cell(1, column).value: column for column in range(1, sheet.max_column + 1)
    }
    required = {
        "semantic_fidelity_1_5",
        "naturalness_1_5",
        "single_correct_answer",
        "native_review_status",
    }
    if not required.issubset(headers):
        raise RuntimeError(f"Missing review fields: {sorted(required - set(headers))}")
    rows_seen = 0
    for row_index in range(2, sheet.max_row + 1):
        if not sheet.cell(row_index, headers["question_id"]).value:
            continue
        rows_seen += 1
        sheet.cell(row_index, headers["semantic_fidelity_1_5"], 4)
        sheet.cell(row_index, headers["naturalness_1_5"], 4)
        sheet.cell(row_index, headers["single_correct_answer"], "Yes")
        sheet.cell(row_index, headers["native_review_status"], "Approved")
    if rows_seen != 960:
        raise RuntimeError(f"Expected 960 workbook MCQ rows, found {rows_seen}")

    readme = workbook["README"]
    target_row = readme.max_row + 1
    readme.cell(target_row, 1, "External review provenance")
    readme.cell(target_row, 2, PROVENANCE)
    for column in (1, 2):
        source = readme.cell(target_row - 1, column)
        target = readme.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        target.number_format = source.number_format
    readme.row_dimensions[target_row].height = 75
    workbook.save(APPROVED)


def verify_approved_workbook() -> None:
    pre = openpyxl.load_workbook(PREAPPROVAL, data_only=False)
    approved = openpyxl.load_workbook(APPROVED, data_only=False)
    if pre.sheetnames != approved.sheetnames:
        raise RuntimeError("Sheet names changed during approval")
    pre_sheet = pre["MCQ_Translations"]
    approved_sheet = approved["MCQ_Translations"]
    headers = {
        approved_sheet.cell(1, column).value: column
        for column in range(1, approved_sheet.max_column + 1)
    }
    mutable_columns = {
        headers["semantic_fidelity_1_5"],
        headers["naturalness_1_5"],
        headers["single_correct_answer"],
        headers["native_review_status"],
    }
    for row in range(1, pre_sheet.max_row + 1):
        for column in range(1, pre_sheet.max_column + 1):
            if row >= 2 and column in mutable_columns:
                continue
            if pre_sheet.cell(row, column).value != approved_sheet.cell(row, column).value:
                raise RuntimeError(
                    f"Unexpected workbook content change at MCQ_Translations!{row},{column}"
                )
    for name in pre.sheetnames:
        if name in {"MCQ_Translations", "README"}:
            continue
        before = pre[name]
        after = approved[name]
        if before.max_row != after.max_row or before.max_column != after.max_column:
            raise RuntimeError(f"Unexpected dimensions changed in {name}")
        for row in range(1, before.max_row + 1):
            for column in range(1, before.max_column + 1):
                if before.cell(row, column).value != after.cell(row, column).value:
                    raise RuntimeError(f"Unexpected content changed in {name}!{row},{column}")
    approved_values = [
        (
            approved_sheet.cell(row, headers["semantic_fidelity_1_5"]).value,
            approved_sheet.cell(row, headers["naturalness_1_5"]).value,
            approved_sheet.cell(row, headers["single_correct_answer"]).value,
            approved_sheet.cell(row, headers["native_review_status"]).value,
        )
        for row in range(2, approved_sheet.max_row + 1)
        if approved_sheet.cell(row, headers["question_id"]).value
    ]
    if len(approved_values) != 960 or set(approved_values) != {(4, 4, "Yes", "Approved")}:
        raise RuntimeError("Approved workbook review values are incomplete or inconsistent")


def update_reports_and_metadata(backup_dir: Path, timestamp: str) -> dict[str, object]:
    design_path = PROJECT / "reports" / "comprehension_design_report.md"
    design_text = design_path.read_text(encoding="utf-8")
    marker = "## External review provenance"
    if marker not in design_text:
        design_text += f"\n{marker}\n\n{PROVENANCE}\n"
    design_path.write_text(design_text, encoding="utf-8")

    hashes = {
        "untouched_original_workbook_sha256": sha256(ORIGINAL),
        "preapproval_completed_workbook_sha256": sha256(PREAPPROVAL),
        "approved_workbook_sha256": sha256(APPROVED),
    }
    metadata = {
        "experiment": "multilingual_comprehension_check",
        "approval_timestamp_utc": timestamp,
        "status": "approved_ready_for_live_validation",
        "review_provenance": PROVENANCE,
        "reviewer_names_recorded": False,
        "approved_multilingual_rows": 960,
        "semantic_fidelity_encoding": 4,
        "naturalness_encoding": 4,
        "single_correct_answer": "Yes",
        "native_review_status": "Approved",
        "backup_directory": str(backup_dir),
        "hashes": hashes,
        "models": [
            "openai/gpt-4o-2024-11-20",
            "anthropic/claude-sonnet-4.6",
            "google/gemini-3.5-flash",
        ],
        "temperature": 0,
        "max_tokens": 24,
        "gemini_reasoning_effort": "minimal",
        "expected_work_units": 720,
        "expected_scored_answers": 2880,
        "paid_api_requests_at_approval": 0,
    }
    (PROJECT / "run_metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (PROCESSED / "approval_metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return metadata


def main() -> None:
    if not PREAPPROVAL.exists() or not CSV_PATH.exists():
        raise FileNotFoundError("Pre-approval workbook or processed multilingual CSV is missing")
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    backup_dir = PROJECT / "data" / "input" / "approval_backups" / timestamp
    backup_dir.mkdir(parents=True, exist_ok=False)
    update_csv(backup_dir)
    update_workbook(backup_dir)
    verify_approved_workbook()
    metadata = update_reports_and_metadata(backup_dir, timestamp)
    print(json.dumps(metadata, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
