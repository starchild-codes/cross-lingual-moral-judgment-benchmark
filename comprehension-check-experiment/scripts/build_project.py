from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import re
import shutil
import sys
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from copy import copy
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import requests
from openpyxl.styles import Alignment, Font, PatternFill


PROJECT = Path(__file__).resolve().parents[1]
REPO = PROJECT.parent
SOURCE_WORKBOOK = Path(
    os.environ.get(
        "COMPREHENSION_SOURCE_WORKBOOK",
        str(PROJECT / "data" / "input" / "multilingual_comprehension_check.xlsx"),
    )
)
ORIGINAL_COPY = PROJECT / "data" / "input" / SOURCE_WORKBOOK.name
COMPLETED_WORKBOOK = PROJECT / "multilingual_comprehension_check_completed.xlsx"
LANGUAGES = {
    "hi": "Hindi",
    "bn": "Bengali",
    "ta": "Tamil",
    "es": "Spanish",
    "ja": "Japanese",
    "ar": "Arabic",
}
MODELS = {
    "chatgpt": "openai/gpt-4o-2024-11-20",
    "claude": "anthropic/claude-sonnet-4.6",
    "gemini_flash": "google/gemini-3.5-flash",
}
FOUNDATIONS = [
    "Care/Harm",
    "Fairness/Cheating",
    "Loyalty/Betrayal",
    "Authority/Subversion",
    "Sanctity/Degradation",
]
QUESTION_TYPES = ["Actor", "Main action", "Affected party/object", "Consequence/outcome"]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv(path: Path, rows: list[dict[str, Any]], headers: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if headers is None:
        headers = list(rows[0]) if rows else []
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def load_scenarios() -> dict[str, dict[str, str]]:
    result: dict[str, dict[str, str]] = {}
    for filename in ("scenarios.csv", "scenarios_extension.csv"):
        for row in read_csv(REPO / "data" / filename):
            scenario_id = row.get("scenario_id") or row.get("scenarios_id")
            if scenario_id:
                if scenario_id.isdigit():
                    scenario_id = f"S{int(scenario_id):02d}"
                result[scenario_id] = row
    if len(result) != 50:
        raise RuntimeError(f"Expected 50 scenarios, found {len(result)}")
    return result


def load_ratings() -> pd.DataFrame:
    frames = [
        pd.read_csv(REPO / "results" / "processed" / "full_1782215308316_vrq93w.ratings.csv"),
        pd.read_csv(REPO / "results" / "processed" / "extension_full_1782729062659_xqetbf.ratings.csv"),
    ]
    data = pd.concat(frames, ignore_index=True)
    data = data[
        (data["modelKey"].isin(MODELS))
        & (data["status"] == "succeeded")
        & data["parsedRating"].notna()
    ].copy()
    duplicates = data.duplicated(["scenarioId", "modelKey", "conditionId"]).sum()
    if duplicates:
        raise RuntimeError(f"Rating source has {duplicates} duplicate model-condition observations")
    return data


def load_human_validation() -> dict[str, dict[str, Any]]:
    path = (
        REPO
        / "results"
        / "processed"
        / "scenario_validation_trial2"
        / "scenario_validation_merged_3coders.csv"
    )
    rows = read_csv(path)
    result: dict[str, dict[str, Any]] = {}
    for row in rows:
        labels = [row.get(f"Coder_{letter}_label", "") for letter in "ABC"]
        majority = row.get("majority_vote_label", "")
        intended = row["intended_mft_foundation"]
        result[row["scenario_id"]] = {
            "intended": intended,
            "majority": majority,
            "majority_resolved": bool(majority),
            "majority_matches_intended": majority == intended,
            "num_coders_matching_intended": sum(label == intended for label in labels),
            "all_three_disagree": row.get("all_three_disagree", "").lower() == "true",
        }
    return result


def compute_shift_audit(
    scenarios: dict[str, dict[str, str]],
    ratings: pd.DataFrame,
    validation: dict[str, dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    audit: list[dict[str, Any]] = []
    effects: dict[str, dict[str, list[float]]] = {}
    for scenario_id in sorted(scenarios):
        block = ratings[ratings["scenarioId"] == scenario_id]
        lookup = {
            (row.modelKey, row.conditionId): float(row.parsedRating)
            for row in block.itertuples()
        }
        input_effects: list[float] = []
        adaptation_effects: list[float] = []
        response_effects: list[float] = []
        all_effects: list[float] = []
        for model in MODELS:
            baseline = lookup[(model, "en_en")]
            for language in LANGUAGES:
                literal_en = lookup[(model, f"{language}_translation_reason_en")]
                literal_native = lookup[(model, f"{language}_translation_reason_{language}")]
                adapted_en = lookup[(model, f"{language}_adapted_reason_en")]
                adapted_native = lookup[(model, f"{language}_adapted_reason_{language}")]
                input_delta = abs(literal_en - baseline)
                adaptation_delta = abs(adapted_en - literal_en)
                response_lit = abs(literal_native - literal_en)
                response_adapted = abs(adapted_native - adapted_en)
                input_effects.append(input_delta)
                adaptation_effects.append(adaptation_delta)
                response_effects.extend([response_lit, response_adapted])
                all_effects.extend([input_delta, adaptation_delta, response_lit, response_adapted])
        effects[scenario_id] = {
            "input": input_effects,
            "adaptation": adaptation_effects,
            "response": response_effects,
            "all": all_effects,
        }
        human = validation[scenario_id]
        audit.append(
            {
                "scenario_id": scenario_id,
                "foundation": human["intended"],
                "human_majority_foundation": human["majority"],
                "majority_resolved": human["majority_resolved"],
                "majority_matches_intended": human["majority_matches_intended"],
                "num_coders_matching_intended": human["num_coders_matching_intended"],
                "all_three_disagree": human["all_three_disagree"],
                "input_language_mad": sum(input_effects) / len(input_effects),
                "adaptation_mad": sum(adaptation_effects) / len(adaptation_effects),
                "response_language_mad": sum(response_effects) / len(response_effects),
                "maximum_observed_shift": max(all_effects),
                "nonzero_input_effect_count": sum(value > 0 for value in input_effects),
                "nonzero_total_effect_count": sum(value > 0 for value in all_effects),
                "eligible": human["majority_resolved"] and human["majority_matches_intended"],
                "rank_within_foundation": "",
                "selection_tier": "",
                "selection_status": "Not selected",
                "selection_rationale": "",
            }
        )

    selected_ids: dict[str, str] = {}
    for foundation in FOUNDATIONS:
        eligible = [row for row in audit if row["foundation"] == foundation and row["eligible"]]
        eligible.sort(
            key=lambda row: (
                -row["input_language_mad"],
                -row["num_coders_matching_intended"],
                row["scenario_id"],
            )
        )
        for rank, row in enumerate(eligible, start=1):
            row["rank_within_foundation"] = rank
        for row in eligible[:2]:
            selected_ids[row["scenario_id"]] = "high"
        low_order = sorted(
            eligible,
            key=lambda row: (
                row["input_language_mad"],
                -row["num_coders_matching_intended"],
                row["scenario_id"],
            ),
        )
        for row in low_order[:2]:
            selected_ids[row["scenario_id"]] = "low"

    for row in audit:
        scenario_id = row["scenario_id"]
        if scenario_id in selected_ids:
            tier = selected_ids[scenario_id]
            row["selection_tier"] = tier
            row["selection_status"] = "Selected"
            edge = "two highest" if tier == "high" else "two lowest"
            row["selection_rationale"] = (
                f"Human-majority-resolved and matched intended foundation; among the {edge} "
                "primary input-language MAD values in this foundation, with coder agreement "
                "and scenario ID used as deterministic tie-breakers."
            )
        elif not row["eligible"]:
            reason = (
                "Excluded: no resolved human majority."
                if not row["majority_resolved"]
                else "Excluded: human majority did not match intended foundation."
            )
            row["selection_rationale"] = reason
        else:
            row["selection_rationale"] = "Eligible but outside the selected high/low extremes."

    selected = [row.copy() for row in audit if row["selection_status"] == "Selected"]
    selected.sort(
        key=lambda row: (
            FOUNDATIONS.index(row["foundation"]),
            0 if row["selection_tier"] == "high" else 1,
            -row["input_language_mad"] if row["selection_tier"] == "high" else row["input_language_mad"],
            row["scenario_id"],
        )
    )
    if len(selected) != 20:
        raise RuntimeError(f"Expected 20 selected scenarios, found {len(selected)}")
    distribution = Counter((row["foundation"], row["selection_tier"]) for row in selected)
    expected = Counter((foundation, tier) for foundation in FOUNDATIONS for tier in ("high", "low"))
    if any(distribution[key] != 2 for key in expected):
        raise RuntimeError(f"Unbalanced scenario selection: {distribution}")
    for row in selected:
        source = scenarios[row["scenario_id"]]
        row["title"] = source.get("title", row["scenario_id"])
        row["scenario_text_en"] = source["text_en"]
        row["selection_summary"] = (
            f"Primary input-language MAD={row['input_language_mad']:.3f}; "
            f"adaptation MAD={row['adaptation_mad']:.3f}; "
            f"response-language MAD={row['response_language_mad']:.3f}; "
            f"maximum shift={row['maximum_observed_shift']:.0f}; "
            f"nonzero effects={row['nonzero_total_effect_count']}/72."
        )
    return audit, selected


BANK: dict[str, dict[str, Any]] = {
    "S26": {
        "descriptor": "a conflict involving weekly assistance",
        "answers": [
            ("The neighbor who usually collected the medicine", ["The elderly neighbor", "The cousin holding the wedding", "A pharmacist"]),
            ("She attended the wedding without arranging the medicine pickup", ["She postponed the wedding", "She delivered the medicine early", "She asked the pharmacist to visit"]),
            ("The elderly neighbor who depended on the weekly pickup", ["The wedding guests", "The pharmacist's family", "A hospital administrator"]),
            ("He went more than a day without his prescribed medication", ["He received a replacement prescription immediately", "The wedding was cancelled", "A relative moved in that evening"]),
        ],
    },
    "S40": {
        "descriptor": "a delayed bus route",
        "answers": [
            ("The bus driver", ["The child", "A passenger", "The route supervisor"]),
            ("The driver continued the route to remain on schedule", ["The driver stopped and waited with the child", "The child boarded another bus", "A passenger called the school"]),
            ("The child left behind at the stop", ["The route supervisor", "A mechanic", "The bus company accountant"]),
            ("The bus continued while the child remained at the stop", ["The route was cancelled", "The child was driven home by the driver", "The bus returned immediately"]),
        ],
    },
    "S05": {
        "descriptor": "missed appointment support",
        "answers": [
            ("The brother who had promised to accompany his sister", ["The anxious sister", "Her doctor", "A professional contact"]),
            ("He cancelled an hour before the appointment to attend a career lunch", ["He attended the appointment and missed lunch", "He rescheduled the appointment himself", "He asked the doctor to visit at home"]),
            ("The younger sister who could not manage the appointment alone", ["The professional contacts", "The hospital receptionist", "Their parents"]),
            ("She cancelled the appointment and her condition later deteriorated", ["She attended alone and recovered", "The lunch was postponed", "The doctor ended her treatment"]),
        ],
    },
    "S02": {
        "descriptor": "an incident beside a park pond",
        "answers": [
            ("The passerby who saw the struggling child", ["The gardener", "The child", "A parent"]),
            ("He walked on because he did not want his clothes to get wet", ["He entered the pond immediately", "He called the gardener before leaving", "He carried the child to a clinic"]),
            ("The child struggling in the shallow pond", ["The gardener's equipment", "The meeting organizer", "A nearby family"]),
            ("A gardener found the child two minutes later, shaken and frightened", ["The child climbed out before anyone arrived", "The passerby missed his meeting", "The pond was permanently closed"]),
        ],
    },
    "S28": {
        "descriptor": "an error during a cash transaction",
        "answers": [
            ("The customer who noticed the mistake", ["The cashier", "A store manager", "The next customer"]),
            ("The customer kept the extra change", ["The customer returned the money", "The cashier deliberately issued a refund", "The manager cancelled the purchase"]),
            ("The money mistakenly given as extra change", ["A damaged product", "A loyalty card", "The customer's receipt"]),
            ("The customer retained money that had been given by mistake", ["The cashier corrected the amount immediately", "The purchase was refunded in full", "The store closed the register"]),
        ],
    },
    "S20": {
        "descriptor": "a disputed hiring decision",
        "answers": [
            ("The senior hiring-panel member", ["The junior panellist", "The passed-over candidate", "The firm's receptionist"]),
            ("She stayed silent when biased comments influenced the decision", ["She stopped the meeting and reopened interviews", "She reported the comments before the decision", "She withdrew from the panel"]),
            ("The strongest candidate who was passed over", ["The firm's clients", "The panel chair's assistant", "A newly hired receptionist"]),
            ("The position was filled and the candidate had no recourse", ["The candidate was hired after an appeal", "The panel repeated every interview", "The junior panellist resigned"]),
        ],
    },
    "S18": {
        "descriptor": "credit for a university report",
        "answers": [
            ("The student who was asked to handle formatting and submission", ["The co-author who wrote half the report", "The course tutor", "A library assistant"]),
            ("He listed himself as primary author and his co-author as a contributor", ["He listed both students equally", "He removed his own name", "He asked the tutor to assign authorship"]),
            ("The co-author whose equal contribution was understated", ["The tutor's research team", "A different class", "The university library"]),
            ("He received the higher grade, and she learned after grades were released", ["Both received identical grades after correcting the form", "The report was never submitted", "The tutor cancelled the assignment"]),
        ],
    },
    "S36": {
        "descriptor": "a workplace promotion",
        "answers": [
            ("The manager making the promotion decision", ["The less qualified relative", "The most deserving employee", "A customer"]),
            ("The manager promoted a less qualified relative", ["The manager ran an anonymous skills test", "The manager promoted the strongest employee", "The manager cancelled the role"]),
            ("The more deserving employee who was not promoted", ["The manager's customer", "A supplier", "The relative's former school"]),
            ("The less qualified relative received the promotion", ["The position remained vacant", "The deserving employee received two promotions", "The manager transferred departments"]),
        ],
    },
    "S37": {
        "descriptor": "confidential information during military service",
        "answers": [
            ("The soldier with access to confidential information", ["The friend living abroad", "The commanding officer", "A journalist"]),
            ("The soldier revealed confidential information to protect a friend", ["The soldier destroyed the information", "The friend reported the soldier", "The commander released the information publicly"]),
            ("The confidentiality of the military information", ["A sports team's schedule", "A school examination", "A family photograph"]),
            ("The information was disclosed for the purpose of protecting the friend", ["The friend was arrested by the soldier", "The information remained completely secret", "The operation was publicly cancelled"]),
        ],
    },
    "S42": {
        "descriptor": "a change in a business partnership",
        "answers": [
            ("The partner who received a more profitable offer", ["The longtime teammate", "A customer", "An accountant"]),
            ("The partner abandoned the teammate to join a competitor", ["The partners rejected every new offer", "The teammate joined the competitor first", "The partners sold the company together"]),
            ("The longtime teammate in the existing partnership", ["The competitor's customers", "A bank manager", "The local government"]),
            ("The existing partnership was left behind for the competitor's offer", ["The competitor withdrew the offer", "The partners expanded together", "The teammate bought the competitor"]),
        ],
    },
    "S06": {
        "descriptor": "confidential information in a promotion competition",
        "answers": [
            ("The colleague competing for the same promotion", ["The employee with the prior dismissal", "The department manager", "A former employer"]),
            ("She disclosed her colleague's confidential past dismissal to the manager", ["She withdrew from the promotion contest", "She defended him to the manager", "She asked him to disclose it himself"]),
            ("The colleague who had shared the information in confidence", ["The department's customers", "A new applicant", "The former employer's accountant"]),
            ("He did not receive the promotion", ["Both colleagues were promoted", "The manager ignored the information and resigned", "The promotion process was cancelled"]),
        ],
    },
    "S47": {
        "descriptor": "confidential plans at a company",
        "answers": [
            ("The employee with access to the plans", ["The competitor", "The company director", "A customer"]),
            ("The employee shared confidential company plans with a competitor", ["The employee encrypted the plans", "The competitor returned the plans unread", "The director published the plans"]),
            ("The company whose plans were confidential", ["A local school", "The employee's family", "A sports club"]),
            ("The competitor received the plans for the employee's personal gain", ["The plans remained inside the company", "The employee donated the plans to a library", "The competitor ceased operating"]),
        ],
    },
    "S48": {
        "descriptor": "an order during a police operation",
        "answers": [
            ("The police officer who received the order", ["The senior officer", "A civilian witness", "A dispatcher"]),
            ("The officer refused to follow a lawful order", ["The officer followed the order immediately", "The senior officer withdrew from the operation", "A witness issued a new order"]),
            ("The operation and its lawful chain of command", ["A private business contract", "A school timetable", "A family celebration"]),
            ("The lawful order was not followed during the operation", ["The operation ended before any order was given", "The order was ruled unlawful", "The senior officer left the police"]),
        ],
    },
    "S43": {
        "descriptor": "an evacuation during an emergency",
        "answers": [
            ("The firefighter who received the evacuation order", ["The incident commander", "A journalist", "A building owner"]),
            ("The firefighter refused the commander's order to evacuate", ["The firefighter evacuated immediately", "The commander cancelled the emergency", "The building owner led the operation"]),
            ("The emergency evacuation directed by incident command", ["A routine training schedule", "A newspaper report", "A property sale"]),
            ("The evacuation order was not followed", ["The order was never issued", "The emergency became a drill", "The commander resigned before arriving"]),
        ],
    },
    "S12": {
        "descriptor": "a hospital protocol discussion",
        "answers": [
            ("The junior doctor who compiled the delay data", ["The supervising consultant", "A nursing administrator", "An elderly patient"]),
            ("She presented the data publicly without first discussing it with her supervisor", ["She destroyed the data", "She discussed it privately and cancelled the meeting", "She asked the supervisor to present it"]),
            ("The supervising consultant and his established intake protocol", ["The hospital cafeteria", "A private pharmacy", "A medical school applicant"]),
            ("The supervisor did not speak to her directly for two weeks", ["The protocol was changed during the meeting", "She was immediately promoted", "The hospital closed the ward"]),
        ],
    },
    "S33": {
        "descriptor": "instructions in a laboratory",
        "answers": [
            ("The laboratory assistant", ["The supervisor", "A visiting student", "A patient"]),
            ("The assistant ignored the supervisor's instructions to save time", ["The supervisor ignored the assistant", "The assistant requested more instructions", "The laboratory stopped all work"]),
            ("The supervisor's required laboratory procedure", ["A cafeteria order", "A travel itinerary", "A university application"]),
            ("The instructions were bypassed in order to finish faster", ["The work took longer because every instruction was followed", "The supervisor cancelled the task", "The laboratory moved buildings"]),
        ],
    },
    "S23": {
        "descriptor": "private conduct in a parish role",
        "answers": [
            ("The parish priest bound by the role's vows", ["The woman in the neighboring town", "A parish member", "The bishop"]),
            ("He conducted a private romantic relationship contrary to his vows", ["He publicly resigned before the relationship", "He reported another priest", "He ended all contact with the neighboring town"]),
            ("The vows and expectations attached to his parish role", ["A commercial lease", "A school examination", "A medical prescription"]),
            ("A parish member discovered the relationship and told no one", ["The congregation held a public protest", "The bishop announced the relationship", "The relationship caused a public scandal"]),
        ],
    },
    "S22": {
        "descriptor": "a family name used for a snack product",
        "answers": [
            ("The family member who ran the food business", ["Her grandfather", "Her parents", "A retail customer"]),
            ("She used her grandfather's sacred name for a snack line without consulting family", ["She closed the food business", "Her parents chose the product name", "She used an unrelated brand name"]),
            ("The family and the revered name of the grandfather", ["A competing restaurant", "A delivery driver", "A local school"]),
            ("Her parents learned about the product when they saw the packaging", ["The product was never packaged", "The grandfather selected the design", "The business was sold before launch"]),
        ],
    },
    "S49": {
        "descriptor": "damage at a cemetery",
        "answers": [
            ("The people vandalizing the cemetery", ["The memorial owners", "A groundskeeper", "A visitor reporting damage"]),
            ("They left trash and damaged memorials", ["They cleaned every memorial", "They planted flowers with permission", "They reported storm damage"]),
            ("The cemetery and its memorials", ["A nearby sports field", "A shopping center", "A school laboratory"]),
            ("The site was left with trash and damaged memorials", ["The cemetery was restored by the same people", "Every memorial was moved indoors", "The groundskeeper closed a parking lot"]),
        ],
    },
    "S24": {
        "descriptor": "inherited ceremonial garments",
        "answers": [
            ("The family member who inherited the garments", ["His mother", "The vintage dealer", "The eventual customer"]),
            ("He sold the garments without telling his mother", ["He returned them to the family trunk", "His mother donated them", "He wore them at a family wedding"]),
            ("The inherited garments and the family connection attached to them", ["A newly purchased suit", "The dealer's shop furniture", "A cousin's vehicle"]),
            ("A customer wore them to a fancy-dress party, and his mother learned from a cousin", ["The garments remained in the family trunk", "The dealer returned them before resale", "His mother approved the sale beforehand"]),
        ],
    },
}


def build_english_questions(selected: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for scenario_index, selection in enumerate(selected):
        scenario_id = selection["scenario_id"]
        bank = BANK[scenario_id]
        for question_index, question_type in enumerate(QUESTION_TYPES):
            answer, distractors = bank["answers"][question_index]
            correct_index = (scenario_index + question_index) % 4
            options = list(distractors)
            options.insert(correct_index, answer)
            question_stem = {
                0: "who performed the central action?",
                1: "what central action was taken?",
                2: "who or what was directly affected?",
                3: "what consequence or outcome is explicitly stated?",
            }[question_index]
            question = f"In the scenario concerning {bank['descriptor']}, {question_stem}"
            row = {
                "question_id": f"{scenario_id}_Q{question_index + 1}",
                "scenario_id": scenario_id,
                "foundation": selection["foundation"],
                "shift_tier": selection["selection_tier"],
                "question_type": question_type,
                "scenario_text_en": selection["scenario_text_en"],
                "question_en": question,
                "option_A": options[0],
                "option_B": options[1],
                "option_C": options[2],
                "option_D": options[3],
                "correct_option": "ABCD"[correct_index],
                "ambiguity_flag": "Pending human review",
                "ambiguity_note": "",
                "reviewer_1_status": "Pending",
                "reviewer_2_status": "Pending",
                "notes": "Factual comprehension item; no moral judgment is requested.",
            }
            rows.append(row)
    if len(rows) != 80:
        raise RuntimeError(f"Expected 80 English questions, found {len(rows)}")
    answer_distribution = Counter(row["correct_option"] for row in rows)
    if answer_distribution != Counter({"A": 20, "B": 20, "C": 20, "D": 20}):
        raise RuntimeError(f"Unbalanced answer positions: {answer_distribution}")
    return rows


def translate_strings(strings: list[str]) -> dict[str, dict[str, str]]:
    cache_path = PROJECT / "data" / "processed" / "translation_cache.json"
    cache: dict[str, dict[str, str]] = {}
    if cache_path.exists():
        cache = json.loads(cache_path.read_text(encoding="utf-8"))
    marker = " ZZZSEPZZZ "
    for language in LANGUAGES:
        language_cache = cache.setdefault(language, {})
        missing = [text for text in strings if text not in language_cache]
        batches: list[list[str]] = []
        current: list[str] = []
        size = 0
        for text in missing:
            addition = len(text) + len(marker)
            if current and size + addition > 3200:
                batches.append(current)
                current = []
                size = 0
            current.append(text)
            size += addition
        if current:
            batches.append(current)
        for index, batch in enumerate(batches, start=1):
            joined = marker.join(batch)
            translated_joined = ""
            for attempt in range(3):
                try:
                    response = requests.get(
                        "https://translate.googleapis.com/translate_a/single",
                        params={"client": "gtx", "sl": "en", "tl": language, "dt": "t", "q": joined},
                        timeout=60,
                    )
                    response.raise_for_status()
                    payload = response.json()
                    translated_joined = "".join(part[0] for part in payload[0])
                    break
                except Exception:
                    if attempt == 2:
                        raise
                    time.sleep(1.5 * (attempt + 1))
            parts = [
                part.strip()
                for part in re.split(r"\s*ZZZSEPZZZ\s*", translated_joined)
                if part.strip()
            ]
            if len(parts) != len(batch):
                def translate_one(source: str) -> str:
                    for attempt in range(5):
                        try:
                            response = requests.get(
                                "https://translate.googleapis.com/translate_a/single",
                                params={
                                    "client": "gtx",
                                    "sl": "en",
                                    "tl": language,
                                    "dt": "t",
                                    "q": source,
                                },
                                timeout=60,
                            )
                            response.raise_for_status()
                            payload = response.json()
                            translated = "".join(part[0] for part in payload[0]).strip()
                            return translated
                        except Exception:
                            if attempt == 4:
                                raise
                            time.sleep(2 * (attempt + 1))
                    raise RuntimeError("Unreachable translation retry state")

                with ThreadPoolExecutor(max_workers=8) as executor:
                    parts = list(executor.map(translate_one, batch))
            for source, translated in zip(batch, parts):
                language_cache[source] = translated.strip()
            print(f"Translated {language} batch {index}/{len(batches)}", flush=True)
            cache_path.write_text(
                json.dumps(cache, ensure_ascii=False, indent=2, sort_keys=True),
                encoding="utf-8",
            )
    return cache


def build_multilingual_questions(
    english: list[dict[str, Any]],
    scenarios: dict[str, dict[str, str]],
) -> list[dict[str, Any]]:
    unique_strings = sorted(
        {
            str(row[column])
            for row in english
            for column in ("question_en", "option_A", "option_B", "option_C", "option_D")
        }
        | {
            "Answer all four questions using only the option letters. Return exactly four letters separated by commas, for example: A,B,C,D. Do not add explanations."
        }
    )
    translations = translate_strings(unique_strings)
    rows: list[dict[str, Any]] = []
    for item in english:
        scenario = scenarios[item["scenario_id"]]
        for language_code, language_name in LANGUAGES.items():
            for version, suffix in (("literal", "b"), ("adapted", "c")):
                row = {
                    "question_id": item["question_id"],
                    "scenario_id": item["scenario_id"],
                    "foundation": item["foundation"],
                    "shift_tier": item["shift_tier"],
                    "language_code": language_code,
                    "language": language_name,
                    "scenario_version": version,
                    "question_type": item["question_type"],
                    "scenario_text_final": scenario[f"text_{language_code}_{suffix}"],
                    "question_final": translations[language_code][item["question_en"]],
                    "option_A": translations[language_code][item["option_A"]],
                    "option_B": translations[language_code][item["option_B"]],
                    "option_C": translations[language_code][item["option_C"]],
                    "option_D": translations[language_code][item["option_D"]],
                    "correct_option": item["correct_option"],
                    "native_review_status": "Pending",
                    "semantic_fidelity_1_5": "",
                    "naturalness_1_5": "",
                    "single_correct_answer": "",
                    "notes": (
                        "Automated machine-translation draft. Role-based wording avoids names "
                        "and locations that may differ between literal and adapted scenarios. "
                        "Native-speaker review is required before any model evaluation."
                    ),
                }
                rows.append(row)
    if len(rows) != 960:
        raise RuntimeError(f"Expected 960 multilingual rows, found {len(rows)}")
    return rows


def build_prompts(
    multilingual: list[dict[str, Any]],
    english: list[dict[str, Any]],
) -> dict[tuple[str, str, str], dict[str, str]]:
    grouped: dict[tuple[str, str, str], list[dict[str, Any]]] = {}
    for row in multilingual:
        key = (row["scenario_id"], row["language_code"], row["scenario_version"])
        grouped.setdefault(key, []).append(row)
    instruction_source = (
        "Answer all four questions using only the option letters. Return exactly four letters "
        "separated by commas, for example: A,B,C,D. Do not add explanations."
    )
    cache = json.loads(
        (PROJECT / "data" / "processed" / "translation_cache.json").read_text(encoding="utf-8")
    )
    prompts: dict[tuple[str, str, str], dict[str, str]] = {}
    for key, items in grouped.items():
        items.sort(key=lambda row: row["question_id"])
        language_code = key[1]
        lines = [items[0]["scenario_text_final"], ""]
        for number, row in enumerate(items, start=1):
            lines.extend(
                [
                    f"{number}. {row['question_final']}",
                    f"A. {row['option_A']}",
                    f"B. {row['option_B']}",
                    f"C. {row['option_C']}",
                    f"D. {row['option_D']}",
                    "",
                ]
            )
        lines.append(cache[language_code][instruction_source])
        prompts[key] = {
            "full_prompt": "\n".join(lines),
            "answer_key": ",".join(row["correct_option"] for row in items),
            "question_ids": ",".join(row["question_id"] for row in items),
        }
    if len(prompts) != 240:
        raise RuntimeError(f"Expected 240 scenario-language-version prompts, found {len(prompts)}")
    return prompts


def build_manifest(
    selected: list[dict[str, Any]],
    prompts: dict[tuple[str, str, str], dict[str, str]],
) -> list[dict[str, Any]]:
    selected_map = {row["scenario_id"]: row for row in selected}
    rows: list[dict[str, Any]] = []
    for scenario_id in sorted(selected_map):
        selection = selected_map[scenario_id]
        for language_code in LANGUAGES:
            for version in ("literal", "adapted"):
                prompt = prompts[(scenario_id, language_code, version)]
                for model_key, model_id in MODELS.items():
                    reasoning_effort = "minimal" if model_key == "gemini_flash" else ""
                    work_unit_id = f"{scenario_id}__{language_code}__{version}__{model_key}"
                    rows.append(
                        {
                            "work_unit_id": work_unit_id,
                            "scenario_id": scenario_id,
                            "foundation": selection["foundation"],
                            "shift_tier": selection["selection_tier"],
                            "model_key": model_key,
                            "model_id": model_id,
                            "input_language": language_code,
                            "scenario_version": version,
                            "expected_questions": 4,
                            "question_ids": prompt["question_ids"],
                            "answer_key": prompt["answer_key"],
                            "temperature": 0,
                            "max_tokens": 24,
                            "reasoning_effort": reasoning_effort,
                            "full_prompt": prompt["full_prompt"],
                            "run_status": "pending",
                            "raw_response": "",
                            "parsed_answers": "",
                            "correct_count_0_4": "",
                            "actor_correct": "",
                            "action_correct": "",
                            "affected_correct": "",
                            "consequence_correct": "",
                            "prompt_tokens": "",
                            "completion_tokens": "",
                            "total_tokens": "",
                            "cost_usd": "",
                            "error_note": "",
                        }
                    )
    if len(rows) != 720:
        raise RuntimeError(f"Expected 720 work units, found {len(rows)}")
    if len({row["work_unit_id"] for row in rows}) != 720:
        raise RuntimeError("Run manifest contains duplicate work-unit IDs")
    return rows


def detect_validation_flags(row: dict[str, Any]) -> tuple[list[str], list[str]]:
    critical: list[str] = []
    review: list[str] = []
    required = [
        "question_id",
        "scenario_id",
        "scenario_text_final",
        "question_final",
        "option_A",
        "option_B",
        "option_C",
        "option_D",
        "correct_option",
    ]
    for column in required:
        if row.get(column, "") == "":
            critical.append(f"missing_{column}")
    if row.get("correct_option") not in "ABCD":
        critical.append("invalid_answer_key")
    options = [row.get(f"option_{letter}", "").strip() for letter in "ABCD"]
    if len(set(options)) != 4:
        critical.append("duplicate_answer_options")
    joined = " ".join([row.get("question_final", "")] + options)
    if "\ufffd" in joined:
        critical.append("malformed_unicode")
    if re.search(r"\b(blameworthy|moral(?:ly)?|ethical|right or wrong)\b", row.get("question_final", ""), re.I):
        critical.append("moral_judgment_language")
    if re.search(r"\b(explain|justify|reasoning|why)\b", row.get("question_final", ""), re.I):
        critical.append("explanation_request")
    if row.get("native_review_status") != "Approved":
        review.append("native_speaker_review_pending")
    if not row.get("semantic_fidelity_1_5"):
        review.append("semantic_fidelity_rating_pending")
    if not row.get("naturalness_1_5"):
        review.append("naturalness_rating_pending")
    if row.get("single_correct_answer") != "Yes":
        review.append("single_correct_answer_review_pending")
    return critical, review


def build_validation_rows(multilingual: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for source in multilingual:
        critical, review = detect_validation_flags(source)
        rows.append(
            {
                "question_id": source["question_id"],
                "scenario_id": source["scenario_id"],
                "language_code": source["language_code"],
                "scenario_version": source["scenario_version"],
                "question_type": source["question_type"],
                "structural_validation": "PASS" if not critical else "FAIL",
                "critical_flag_count": len(critical),
                "critical_flags": ";".join(critical),
                "human_review_required": "Yes" if review else "No",
                "review_items": ";".join(review),
                "native_review_status": source["native_review_status"],
            }
        )
    return rows


def copy_cell_style(source: openpyxl.cell.Cell, target: openpyxl.cell.Cell) -> None:
    if source.has_style:
        target._style = copy(source._style)
    target.number_format = source.number_format
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def write_sheet_rows(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    rows: list[dict[str, Any]],
    headers: list[str],
) -> None:
    existing_max_row = ws.max_row
    existing_max_col = ws.max_column
    style_header = ws.cell(1, min(existing_max_col, 1))
    style_body = ws.cell(min(existing_max_row, 2), min(existing_max_col, 1))
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(1, column, header)
        copy_cell_style(style_header, cell)
    max_needed = max(existing_max_row, len(rows) + 1)
    for row_index in range(2, max_needed + 1):
        source = rows[row_index - 2] if row_index - 2 < len(rows) else None
        for column, header in enumerate(headers, start=1):
            cell = ws.cell(row_index, column)
            if column > existing_max_col or row_index > existing_max_row:
                copy_cell_style(style_body, cell)
            cell.value = source.get(header, "") if source else None
    ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}{len(rows) + 1}"
    for table_name in ws.tables:
        table = ws.tables[table_name]
        start_cell, end_cell = table.ref.split(":")
        min_column = openpyxl.utils.column_index_from_string(
            re.match(r"[A-Z]+", start_cell).group()
        )
        max_column = openpyxl.utils.column_index_from_string(
            re.match(r"[A-Z]+", end_cell).group()
        )
        for offset, table_column in enumerate(table.tableColumns):
            column = min_column + offset
            if column <= max_column:
                table_column.name = str(ws.cell(1, column).value)


def update_workbook(
    selected: list[dict[str, Any]],
    audit: list[dict[str, Any]],
    english: list[dict[str, Any]],
    multilingual: list[dict[str, Any]],
    manifest: list[dict[str, Any]],
) -> None:
    workbook = openpyxl.load_workbook(ORIGINAL_COPY)
    expected_sheets = [
        "README",
        "Selected_20",
        "English_MCQ_Authoring",
        "MCQ_Translations",
        "Run_Manifest_720",
        "Analysis_Plan",
    ]
    if workbook.sheetnames != expected_sheets:
        raise RuntimeError(f"Unexpected source sheets: {workbook.sheetnames}")

    readme = workbook["README"]
    readme["A1"] = "Multilingual Comprehension Check: Completed Preparation Workbook"
    readme["A2"] = (
        "This workbook contains the fully prepared 20-scenario comprehension-check design. "
        "No paid model calls have been made."
    )
    readme["A3"] = (
        "The 960 non-English MCQ rows are automated translation drafts and remain at the "
        "required native-speaker review checkpoint."
    )
    readme["A4"] = (
        "Do not run the smoke test or full evaluation until every native_review_status is "
        "Approved and all fidelity, naturalness, and single-answer review fields are complete."
    )
    readme["A5"] = (
        "Selection uses the mean absolute deviation from the English baseline for literal "
        "non-English input with English response, averaged over 3 models x 6 languages."
    )
    readme["A6"] = (
        "Human-validation safeguard: only resolved human-majority scenarios whose majority "
        "foundation matches the intended foundation were eligible."
    )
    readme["A7"] = "Completed workbook generated without paid API calls."

    selected_headers = [
        "scenario_id",
        "intended_foundation",
        "scenario_text_en",
        "majority_vote_label",
        "majority_matches_intended",
        "coders_matching_intended",
        "shift_metric_abs",
        "shift_tier",
        "selection_status",
        "selection_notes",
        "title",
        "rank_within_foundation",
        "input_language_mad",
        "adaptation_mad",
        "response_language_mad",
        "maximum_observed_shift",
        "nonzero_input_effect_count",
        "nonzero_total_effect_count",
        "human_majority_foundation",
        "num_coders_matching_intended",
        "selection_rationale",
    ]
    selected_rows = [
        {
            **row,
            "intended_foundation": row["foundation"],
            "majority_vote_label": row["human_majority_foundation"],
            "coders_matching_intended": row["num_coders_matching_intended"],
            "shift_metric_abs": row["input_language_mad"],
            "shift_tier": (
                "High shift" if row["selection_tier"] == "high" else "Low/near-zero shift"
            ),
            "selection_status": "Final selected",
            "selection_notes": row["selection_summary"],
        }
        for row in selected
    ]
    write_sheet_rows(workbook["Selected_20"], selected_rows, selected_headers)

    english_headers = [
        "question_id",
        "scenario_id",
        "foundation",
        "shift_tier",
        "question_type",
        "scenario_text_en",
        "question_en",
        "option_A_en",
        "option_B_en",
        "option_C_en",
        "option_D_en",
        "correct_option",
        "reviewer_1_status",
        "reviewer_2_status",
        "ambiguity_flag",
        "author_notes",
    ]
    english_workbook_rows = [
        {
            **row,
            "shift_tier": "High shift" if row["shift_tier"] == "high" else "Low/near-zero shift",
            "option_A_en": row["option_A"],
            "option_B_en": row["option_B"],
            "option_C_en": row["option_C"],
            "option_D_en": row["option_D"],
            "ambiguity_flag": "Unsure",
            "author_notes": row["notes"],
        }
        for row in english
    ]
    write_sheet_rows(workbook["English_MCQ_Authoring"], english_workbook_rows, english_headers)

    multilingual_headers = [
        "question_id",
        "scenario_id",
        "foundation",
        "language_code",
        "language",
        "scenario_version",
        "question_type",
        "scenario_text_final",
        "question_final",
        "option_A",
        "option_B",
        "option_C",
        "option_D",
        "correct_option",
        "native_review_status",
        "semantic_fidelity_1_5",
        "naturalness_1_5",
        "single_correct_answer",
        "notes",
        "shift_tier",
    ]
    write_sheet_rows(workbook["MCQ_Translations"], multilingual, multilingual_headers)

    manifest_headers = [
        "work_unit_id",
        "scenario_id",
        "foundation",
        "shift_tier",
        "model_key",
        "input_language",
        "scenario_version",
        "expected_questions",
        "run_status",
        "raw_response",
        "parsed_answers",
        "correct_count_0_4",
        "actor_correct",
        "action_correct",
        "affected_correct",
        "consequence_correct",
        "error_note",
        "model_id",
        "question_ids",
        "answer_key",
        "temperature",
        "max_tokens",
        "reasoning_effort",
        "full_prompt",
        "prompt_tokens",
        "completion_tokens",
        "total_tokens",
        "cost_usd",
    ]
    manifest_workbook_rows = [
        {
            **row,
            "shift_tier": "High shift" if row["shift_tier"] == "high" else "Low/near-zero shift",
            "run_status": "Not started",
        }
        for row in manifest
    ]
    write_sheet_rows(workbook["Run_Manifest_720"], manifest_workbook_rows, manifest_headers)

    if "Scenario_Shift_Audit" in workbook.sheetnames:
        del workbook["Scenario_Shift_Audit"]
    audit_ws = workbook.create_sheet("Scenario_Shift_Audit")
    audit_headers = list(audit[0])
    write_sheet_rows(audit_ws, audit, audit_headers)
    audit_ws.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in audit_ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column in range(1, len(audit_headers) + 1):
        audit_ws.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 18

    analysis = workbook["Analysis_Plan"]
    start = analysis.max_row + 2
    additions = [
        ("Human-review checkpoint", "All 960 multilingual rows must be approved before model evaluation."),
        ("Primary analysis", "Item accuracy and scenario-level 4/4 comprehension by model, language, version, foundation, and shift tier."),
        ("Failure analysis", "Each missed item will retain question type, selected option, answer key, raw response, and scenario metadata."),
        ("No paid calls", "Preparation and dry-run validation only; smoke and full runs remain blocked."),
    ]
    for offset, (label, value) in enumerate(additions):
        analysis.cell(start + offset, 1, label)
        analysis.cell(start + offset, 2, value)

    if workbook.calculation is not None:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    workbook.save(COMPLETED_WORKBOOK)


def verify_outputs(
    original_hash: str,
    selected: list[dict[str, Any]],
    english: list[dict[str, Any]],
    multilingual: list[dict[str, Any]],
    manifest: list[dict[str, Any]],
) -> dict[str, Any]:
    reopened = openpyxl.load_workbook(COMPLETED_WORKBOOK, data_only=False)
    expected_sheets = {
        "README",
        "Selected_20",
        "English_MCQ_Authoring",
        "MCQ_Translations",
        "Run_Manifest_720",
        "Analysis_Plan",
        "Scenario_Shift_Audit",
    }
    if set(reopened.sheetnames) != expected_sheets:
        raise RuntimeError(f"Completed workbook sheet mismatch: {reopened.sheetnames}")
    if sha256(SOURCE_WORKBOOK) != original_hash or sha256(ORIGINAL_COPY) != original_hash:
        raise RuntimeError("Original workbook hash changed")
    if len(selected) != 20 or len(english) != 80 or len(multilingual) != 960 or len(manifest) != 720:
        raise RuntimeError("Output row-count verification failed")
    for row in english:
        for column in ("question_id", "scenario_id", "option_A", "option_B", "option_C", "option_D", "correct_option"):
            if not row.get(column):
                raise RuntimeError(f"English MCQ missing {column}: {row}")
    for row in multilingual:
        for column in ("question_id", "scenario_id", "option_A", "option_B", "option_C", "option_D", "correct_option"):
            if not row.get(column):
                raise RuntimeError(f"Multilingual MCQ missing {column}: {row}")
    if len({row["work_unit_id"] for row in manifest}) != 720:
        raise RuntimeError("Manifest uniqueness verification failed")
    return {
        "original_workbook_sha256": original_hash,
        "completed_workbook_sha256": sha256(COMPLETED_WORKBOOK),
        "completed_workbook_readable": True,
        "expected_sheets_present": True,
        "selected_scenarios": len(selected),
        "english_mcq_rows": len(english),
        "multilingual_mcq_rows": len(multilingual),
        "unique_manifest_rows": len({row["work_unit_id"] for row in manifest}),
        "paid_api_requests": 0,
    }


def main() -> None:
    if not SOURCE_WORKBOOK.exists():
        raise FileNotFoundError(SOURCE_WORKBOOK)
    PROJECT.joinpath("data", "input").mkdir(parents=True, exist_ok=True)
    original_hash = sha256(SOURCE_WORKBOOK)
    if ORIGINAL_COPY.exists() and sha256(ORIGINAL_COPY) != original_hash:
        raise RuntimeError("Existing original-copy hash does not match the attached workbook")
    if not ORIGINAL_COPY.exists():
        shutil.copy2(SOURCE_WORKBOOK, ORIGINAL_COPY)
    if sha256(ORIGINAL_COPY) != original_hash:
        raise RuntimeError("Original workbook copy failed hash verification")

    scenarios = load_scenarios()
    ratings = load_ratings()
    human_validation = load_human_validation()
    audit, selected = compute_shift_audit(scenarios, ratings, human_validation)
    english = build_english_questions(selected)
    multilingual = build_multilingual_questions(english, scenarios)
    prompts = build_prompts(multilingual, english)
    manifest = build_manifest(selected, prompts)
    validation_rows = build_validation_rows(multilingual)

    processed = PROJECT / "data" / "processed"
    write_csv(processed / "selected_scenarios.csv", selected)
    write_csv(processed / "scenario_shift_audit.csv", audit)
    write_csv(processed / "english_comprehension_questions.csv", english)
    write_csv(processed / "multilingual_comprehension_questions.csv", multilingual)
    write_csv(processed / "run_manifest_720.csv", manifest)
    write_csv(processed / "comprehension_item_validation.csv", validation_rows)
    update_workbook(selected, audit, english, multilingual, manifest)
    verification = verify_outputs(original_hash, selected, english, multilingual, manifest)
    (processed / "build_verification.json").write_text(
        json.dumps(verification, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(verification, indent=2))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
