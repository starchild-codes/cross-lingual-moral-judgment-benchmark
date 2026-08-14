from __future__ import annotations

import csv
import hashlib
import json
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .common import DATABASE_PATH, INPUT_WORKBOOK, RESULTS_DIR, ROOT, configure_utf8_console


BACKUP_WORKBOOK = ROOT / "data" / "input" / "neutral_controls_multilingual_preapproval.xlsx"
VERIFICATION_JSON = RESULTS_DIR / "model_verification.json"
RUN_METADATA_JSON = RESULTS_DIR / "run_metadata.json"
COST_LEDGER_CSV = RESULTS_DIR / "cost_ledger.csv"
EXPECTED_MODELS = {
    "chatgpt": "openai/gpt-4o-2024-11-20",
    "claude": "anthropic/claude-sonnet-4.6",
    "gemini_flash": "google/gemini-3.5-flash",
}
SMOKE_KEYS = {f"C01::{model_key}::en_en" for model_key in EXPECTED_MODELS}
OUTPUT_FILES = [
    "data/processed/neutral_control_conditions.csv",
    "data/results/neutral_control_results.sqlite3",
    "data/results/neutral_control_ratings_raw.csv",
    "data/results/neutral_control_ratings_clean.csv",
    "data/results/neutral_control_diagnostics.csv",
    "data/results/neutral_control_summary.csv",
    "data/results/neutral_control_contrasts.csv",
    "data/results/neutral_control_summary.md",
    "data/results/model_verification.json",
    "data/results/run_metadata.json",
    "data/results/cost_ledger.csv",
    "figures/neutral_control_rating_distribution.png",
    "figures/neutral_control_rating_distribution_plot_data.csv",
    "figures/neutral_control_effects.png",
    "figures/neutral_control_effects_plot_data.csv",
    "reports/neutral_control_report.md",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def headers(sheet) -> dict[str, int]:
    return {
        str(sheet.cell(1, column).value): column
        for column in range(1, sheet.max_column + 1)
    }


def text_digest(path: Path) -> str:
    workbook = load_workbook(path, read_only=False, data_only=False)
    sheet = workbook["Controls_Long"]
    index = headers(sheet)
    digest = hashlib.sha256()
    for row in range(2, sheet.max_row + 1):
        for column in ("scenario_text", "question"):
            value = sheet.cell(row, index[column]).value
            digest.update(str(value).encode("utf-8"))
            digest.update(b"\x00")
    workbook.close()
    return digest.hexdigest().upper()


def verification_cost(records: list[dict[str, Any]]) -> float:
    return sum(
        float(record.get("response_json", {}).get("usage", {}).get("cost") or 0)
        for record in records
    )


def main() -> None:
    configure_utf8_console()
    archive_manifests = sorted(
        (RESULTS_DIR / "archive").glob("*_precorrection_smoke/archive_manifest.json")
    )
    if not archive_manifests:
        raise FileNotFoundError("Pre-correction smoke archive manifest was not found")
    archive_manifest_path = archive_manifests[-1]
    archive_manifest = json.loads(archive_manifest_path.read_text(encoding="utf-8"))
    verification = json.loads(VERIFICATION_JSON.read_text(encoding="utf-8"))

    with sqlite3.connect(DATABASE_PATH) as connection:
        connection.row_factory = sqlite3.Row
        responses = [dict(row) for row in connection.execute("SELECT * FROM responses")]
        attempts = [dict(row) for row in connection.execute("SELECT * FROM attempts")]

    if len(responses) != 375:
        raise RuntimeError(f"Expected 375 active responses; found {len(responses)}")
    unique_keys = {row["unique_key"] for row in responses}
    if len(unique_keys) != 375:
        raise RuntimeError("Active responses contain duplicate work units")
    if not all(row["is_valid"] == 1 for row in responses):
        raise RuntimeError("Active responses contain an invalid final row")

    model_ids = {
        model_key: sorted(
            {
                row["openrouter_model_string"]
                for row in responses
                if row["model_key"] == model_key
            }
        )
        for model_key in EXPECTED_MODELS
    }
    for model_key, expected_id in EXPECTED_MODELS.items():
        if model_ids[model_key] != [expected_id]:
            raise RuntimeError(f"Unexpected model ID for {model_key}: {model_ids[model_key]}")

    parameter_sets = {
        model_key: {
            "temperature": sorted(
                {
                    float(row["request_temperature"])
                    for row in responses
                    if row["model_key"] == model_key
                }
            ),
            "max_tokens": sorted(
                {
                    int(row["request_max_tokens"])
                    for row in responses
                    if row["model_key"] == model_key
                }
            ),
            "reasoning_effort": sorted(
                {
                    row["reasoning_effort"] or None
                    for row in responses
                    if row["model_key"] == model_key
                },
                key=lambda value: "" if value is None else value,
            ),
        }
        for model_key in EXPECTED_MODELS
    }
    expected_effort = {
        "chatgpt": [None],
        "claude": [None],
        "gemini_flash": ["minimal"],
    }
    for model_key, parameters in parameter_sets.items():
        if parameters["temperature"] != [0.0] or parameters["max_tokens"] != [16]:
            raise RuntimeError(f"Inconsistent request parameters for {model_key}: {parameters}")
        if parameters["reasoning_effort"] != expected_effort[model_key]:
            raise RuntimeError(f"Unexpected reasoning effort for {model_key}: {parameters}")

    attempt_parameter_sets = {}
    for model_key in EXPECTED_MODELS:
        model_attempts = [
            row for row in attempts if row["unique_key"].split("::")[1] == model_key
        ]
        attempt_parameter_sets[model_key] = {
            "temperature": sorted({float(row["request_temperature"]) for row in model_attempts}),
            "max_tokens": sorted({int(row["request_max_tokens"]) for row in model_attempts}),
            "reasoning_effort": sorted(
                {row["reasoning_effort"] or None for row in model_attempts},
                key=lambda value: "" if value is None else value,
            ),
        }
        if attempt_parameter_sets[model_key] != parameter_sets[model_key]:
            raise RuntimeError(
                f"Attempt parameters differ from final-row parameters for {model_key}"
            )

    verification_parameter_sets = {}
    for model_key in EXPECTED_MODELS:
        model_verifications = [
            row for row in verification if row["model_key"] == model_key
        ]
        verification_parameter_sets[model_key] = {
            "temperature": sorted(
                {float(row["request_temperature"]) for row in model_verifications}
            ),
            "max_tokens": sorted(
                {int(row["request_max_tokens"]) for row in model_verifications}
            ),
            "reasoning_effort": sorted(
                {row["reasoning_effort"] or None for row in model_verifications},
                key=lambda value: "" if value is None else value,
            ),
        }
        if verification_parameter_sets[model_key] != parameter_sets[model_key]:
            raise RuntimeError(
                f"Verification parameters differ from response parameters for {model_key}"
            )

    active_response_cost = sum(float(row["reported_cost"] or 0) for row in responses)
    replacement_smoke_cost = sum(
        float(row["reported_cost"] or 0)
        for row in responses
        if row["unique_key"] in SMOKE_KEYS
    )
    post_smoke_response_cost = active_response_cost - replacement_smoke_cost
    active_verification_cost = verification_cost(verification)
    archived_response_cost = float(archive_manifest["smoke_response_reported_cost"])
    archived_verification_cost = float(archive_manifest["verification_reported_cost"])
    archived_total_cost = float(archive_manifest["archived_total_reported_cost"])
    full_experiment_cost = post_smoke_response_cost + active_verification_cost
    combined_total_cost = archived_total_cost + replacement_smoke_cost + full_experiment_cost

    malformed_attempts = [
        row for row in attempts if row["http_status"] == 200 and row["is_valid"] == 0
    ]
    retry_attempts = sum(max(int(row["attempt_count"]) - 1, 0) for row in responses)
    active_max_tokens = any(
        "MAX_TOKENS" in str(row.get("raw_output") or "").upper() for row in responses
    ) or any(
        "MAX_TOKENS" in str(row.get("response_json") or "").upper() for row in attempts
    )
    original_text_hash = text_digest(BACKUP_WORKBOOK)
    final_text_hash = text_digest(INPUT_WORKBOOK)
    if original_text_hash != final_text_hash:
        raise RuntimeError("Scenario or question text differs from the preapproval workbook")

    costs = [
        {
            "cost_component": "archived_first_smoke_responses",
            "reported_cost_usd": archived_response_cost,
            "included_in_combined_total": True,
            "note": "Three archived pre-correction smoke responses.",
        },
        {
            "cost_component": "archived_first_smoke_verification",
            "reported_cost_usd": archived_verification_cost,
            "included_in_combined_total": True,
            "note": "Archived pre-correction model-verification requests.",
        },
        {
            "cost_component": "successful_replacement_smoke",
            "reported_cost_usd": replacement_smoke_cost,
            "included_in_combined_total": True,
            "note": "Three valid smoke rows retained as part of the final 375-row dataset.",
        },
        {
            "cost_component": "full_experiment_post_smoke_responses",
            "reported_cost_usd": post_smoke_response_cost,
            "included_in_combined_total": True,
            "note": "The remaining 372 valid response rows.",
        },
        {
            "cost_component": "corrected_full_run_verification",
            "reported_cost_usd": active_verification_cost,
            "included_in_combined_total": True,
            "note": "Exact-ID verification requests, including the resume verification.",
        },
        {
            "cost_component": "combined_total",
            "reported_cost_usd": combined_total_cost,
            "included_in_combined_total": False,
            "note": "Sum of all paid requests above; no component is counted twice.",
        },
    ]
    with COST_LEDGER_CSV.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(costs[0]))
        writer.writeheader()
        writer.writerows(
            {
                **row,
                "reported_cost_usd": f"{row['reported_cost_usd']:.9f}",
            }
            for row in costs
        )

    metadata = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "run_started_at_utc": min(row["created_at"] for row in attempts),
        "run_finished_at_utc": max(row["created_at"] for row in attempts),
        "resumed_after_interruption": True,
        "expected_condition_rows": 125,
        "actual_condition_rows": 125,
        "expected_work_units": 375,
        "actual_unique_valid_rows": 375,
        "active_attempt_rows": len(attempts),
        "retry_attempts": retry_attempts,
        "malformed_attempt_outputs": len(malformed_attempts),
        "unresolved_failures": 0,
        "duplicate_active_work_units": 0,
        "active_output_contains_max_tokens": active_max_tokens,
        "raw_output_stored_for_all_final_rows": all(
            row["raw_output"] is not None for row in responses
        ),
        "token_metadata_rows": sum(row["total_tokens"] is not None for row in responses),
        "cost_metadata_rows": sum(row["reported_cost"] is not None for row in responses),
        "model_ids": model_ids,
        "request_parameters": parameter_sets,
        "attempt_request_parameters": attempt_parameter_sets,
        "verification_request_parameters": verification_parameter_sets,
        "verification_attempts": len(verification),
        "verification_runs": len({row["verification_run"] for row in verification}),
        "costs_usd": {
            "archived_first_smoke_responses": archived_response_cost,
            "archived_first_smoke_verification": archived_verification_cost,
            "archived_first_smoke_total": archived_total_cost,
            "successful_replacement_smoke": replacement_smoke_cost,
            "final_375_response_dataset_including_replacement_smoke": active_response_cost,
            "full_experiment_post_smoke_responses": post_smoke_response_cost,
            "corrected_full_run_verification": active_verification_cost,
            "full_experiment_after_smoke_including_verification": full_experiment_cost,
            "combined_total_all_paid_requests": combined_total_cost,
        },
        "protocol_deviation": (
            "Gemini 3.5 Flash was configured with minimal reasoning effort because its default "
            "reasoning process exhausted the original five-token completion limit before "
            "returning a rating. The completion limit was increased uniformly to 16 tokens "
            "across all evaluated models. No model, prompt, temperature, scenario, question, "
            "or rating-scale changes were made."
        ),
        "workbook": {
            "approved_sha256": sha256(INPUT_WORKBOOK),
            "preapproval_backup_sha256": sha256(BACKUP_WORKBOOK),
            "scenario_question_sha256": final_text_hash,
            "scenario_question_unchanged": True,
        },
        "archive": {
            "folder": str(archive_manifest_path.parent.relative_to(ROOT)),
            "manifest_sha256": sha256(archive_manifest_path),
            "old_smoke_excluded_from_active_store": True,
        },
        "commands": [
            "py -3 -m src.validate_input",
            "py -3 -m src.build_conditions",
            "py -3 -m pytest",
            "py -3 -m src.run_experiment --dry-run",
            "py -3 -m src.run_experiment --smoke-test",
            "py -3 -m src.run_experiment --full",
            "py -3 -m src.run_experiment --resume",
            "py -3 -m src.diagnostics",
            "py -3 -m src.analyze_controls",
            "py -3 -m src.make_figures",
            "py -3 -m src.finalize_run_metadata",
        ],
        "output_files": OUTPUT_FILES,
        "limitations": [
            "All 375 ratings were at the floor value of 1, limiting sensitivity to downward shifts.",
            "Bootstrap contrasts contain only five controls per model-language cell.",
            "The original moral-rating dataset was not used, so no moral-versus-neutral interaction was estimated.",
        ],
    }
    RUN_METADATA_JSON.write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print(f"Wrote run metadata to {RUN_METADATA_JSON}")
    print(f"Wrote cost ledger to {COST_LEDGER_CSV}")
    print(f"combined_total_cost={combined_total_cost:.9f}")


if __name__ == "__main__":
    main()
